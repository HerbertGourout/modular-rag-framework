from __future__ import annotations

from pathlib import Path

from modular_rag.core.models.document import Document


class DocxParser:
    """Extract text from Word (.docx) files using python-docx. Requires [v1] extras."""

    SUPPORTED_EXTENSIONS = {".docx"}

    def supports(self, path: str | Path) -> bool:
        return Path(path).suffix.lower() in self.SUPPORTED_EXTENSIONS

    def parse(self, path: str | Path) -> Document:
        try:
            import docx
        except ImportError as exc:
            raise ImportError(
                "Install 'python-docx' (pip install modular-rag[v1]) to parse Word documents."
            ) from exc

        p = Path(path)
        word_doc = docx.Document(str(p))
        paragraphs = [para.text for para in word_doc.paragraphs if para.text.strip()]
        content = "\n\n".join(paragraphs)
        return Document(
            source=str(p),
            content=content,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            metadata={"paragraph_count": len(paragraphs)},
        )
